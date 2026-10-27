"""
Excel -> Lua 配置表导出器（ExcelToLua.cs 的 Python 移植）

用途：在没有 Unity 编辑器的情况下，用命令行把 `<Mod>/Configs/*.xlsx` 生成
`<Mod>/Configs/Lua/*Config.lua`，产物与 Unity 里点击“生成配置表”一致。

对应实现：jyx2/Assets/Scripts/Utils/Tools/ExcelToLua.cs
    ExportAllLuaFile / ExportSingleLuaFile / GenLuaFile

依赖：openpyxl（pip install openpyxl）

表格格式（详见 jyx2/Assets/Mods/SAMPLE/Configs/README.md）：
    第1行：A1 必须是 "LuaConfigGen"，B1 是配置表名（如 Character）
    第2行：字段中文注释（不参与生成）
    第3行：字段类型 L_number / L_string / L_boolean / 复合(L_number,L_string) / 列表(加 [])
    第4行：字段名（空或以 # 开头的列会被忽略）
    第5行起：数据行（首列为空或以 # 开头的行会被忽略）
"""

from __future__ import annotations

import datetime
import os
import sys

try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

LUA_HEADER = (
    "--[[\n"
    "本文件由编辑器自动生成，如需修改请先修改Excel表格后再使用Unity生成本文件\n"
    "\n"
    "金庸群侠传3D重制版\n"
    "https://github.com/jynew/jynew\n"
    "\n"
    "这是本开源项目文件头，所有代码均使用MIT协议。\n"
    "但游戏内资源和第三方插件、dll等请仔细阅读LICENSE相关授权协议文档。\n"
    "\n"
    "金庸老先生千古！\n"
    "]]\n"
)

STRUCT_SPLIT = ","
ARRAY_SPLIT = "|"
NAME_SPLIT = "-"

L_BOOLEAN = "l_boolean"
L_NUMBER = "l_number"
L_STRING = "l_string"
L_STRUCT = "l_struct"
LUA_TYPES = (L_BOOLEAN, L_NUMBER, L_STRING, L_STRUCT)


class ColDesc:
    """配置表列的描述，对应 C# 的 ColDesc"""

    __slots__ = ("index", "comment", "type_str", "name", "sub_name", "type", "sub_type", "is_array")

    def __init__(self, index: int, name: str, comment: str, type_str: str):
        self.index = index
        self.comment = comment
        self.type_str = type_str
        self.name = name
        self.sub_name: list[str] = []
        self.type: str = L_STRING
        self.sub_type: list[str] = []
        self.is_array: bool = False


def cell_to_str(value) -> str:
    """模拟 .NET DataRow[i].ToString()（ExcelDataReader 读出的类型）"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "True" if value else "False"
    if isinstance(value, int):
        return str(value)
    if isinstance(value, float):
        if value == int(value) and abs(value) < 1e15:
            return str(int(value))
        return repr(value)
    if isinstance(value, datetime.datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S")
    if isinstance(value, datetime.date):
        return value.strftime("%Y-%m-%d 00:00:00")
    if isinstance(value, datetime.time):
        return value.strftime("0001-01-01 %H:%M:%S")
    return str(value)


def string_to_boolean(value: str) -> bool:
    """对应 C# StringToBoolean：空 / 非 0 非 false 即为真"""
    v = value.strip().lower()
    if not v:
        return True
    if v == "false":
        return False
    try:
        return int(v) != 0
    except ValueError:
        return True


def get_col_desc(rows: list[list[str]], col_num: int) -> list[ColDesc]:
    """解析表头，对应 C# ExcelToLua.GetColDesc"""
    col_desc_list: list[ColDesc] = []
    for i in range(col_num):
        name = rows[3][i].strip() if len(rows) > 3 else ""
        if name.startswith("#") or not name:
            continue
        comment = rows[1][i].strip() if len(rows) > 1 else ""
        type_str = rows[2][i].strip().lower() if len(rows) > 2 else ""

        col = ColDesc(i, name, comment, type_str)
        col.is_array = "[]" in type_str
        type_str = type_str.replace("[]", "")

        try:
            type_list = [t for t in type_str.split(STRUCT_SPLIT) if t != ""]
            if len(type_list) > 1:
                col.type = L_STRUCT
                tmp_name = [t for t in name.split(NAME_SPLIT) if t != ""]
                if len(tmp_name) > 1:
                    col.name = tmp_name[0]
                    col.sub_name = [t for t in tmp_name[1].split(STRUCT_SPLIT) if t != ""]
                for t in type_list:
                    if t not in LUA_TYPES:
                        raise ValueError(t)
                    col.sub_type.append(t)
            else:
                if type_str not in LUA_TYPES:
                    raise ValueError(type_str)
                col.type = type_str
        except ValueError:
            # C# 里 Enum.Parse 失败会被 catch，按字符串处理
            print(f"[warn] Lua field type '{col.type_str}' not found, at (3,{i + 1})", file=sys.stderr)
            col.type = L_STRING
            col.sub_type = []
        col_desc_list.append(col)
    return col_desc_list


def get_lua_value(col: ColDesc, value: str) -> str:
    """把单元格/子值转换成 Lua 字面量，对应 C# GetLuaValue"""
    if col.type == L_STRING:
        return '""' if not value.strip() else f"[[{value}]]"
    if col.type == L_BOOLEAN:
        return "true" if string_to_boolean(value) else "false"
    if col.type == L_NUMBER:
        return "nil" if value == "" else value

    # 复合型（struct）
    v_list = value.split(STRUCT_SPLIT)
    parts: list[str] = []
    for idx, v in enumerate(v_list):
        t = col.sub_type[idx] if idx < len(col.sub_type) else L_NUMBER
        if t == L_STRING:
            parts.append('""' if not v.strip() else f"[[{v}]]")
        elif t == L_BOOLEAN:
            parts.append("true" if string_to_boolean(v) else "false")
        else:
            parts.append("nil" if v == "" else v)
    body = ",".join(parts)
    return "{" + body + "}" if len(v_list) > 1 else body




def gen_lua_file(rows: list[list[str]], col_num: int, row_num: int) -> str:
    """生成整份 Lua 文本，对应 C# GenLuaFile"""
    config_type = rows[0][1] if col_num > 1 else ""
    col_desc_list = get_col_desc(rows, col_num)

    sb: list[str] = [LUA_HEADER]
    if not col_desc_list:
        return "".join(sb)

    sb.append("local fieldIdx = {}\n")
    # C# 里 fieldIdx 的编号是“有效列”的顺序（从 1 开始），不是 Excel 的列号
    for i, col in enumerate(col_desc_list):
        sb.append(f"fieldIdx.{col.name} = {i + 1}\n")

    sb.append("local data = {\n")
    start_row_idx = 4
    for i in range(start_row_idx, row_num):
        first_element = rows[i][0].strip() if len(rows[i]) > 0 else ""
        if first_element.startswith("#") or not first_element:
            continue

        one_row = ["{"]
        for j, cur_col in enumerate(col_desc_list):
            raw = rows[i][cur_col.index] if cur_col.index < len(rows[i]) else ""
            content = raw.strip()
            if not cur_col.is_array:
                one_row.append(get_lua_value(cur_col, content))
            else:
                tmp = [x for x in content.split(ARRAY_SPLIT) if x != ""]
                one_row.append("{")
                one_row.append(",".join(get_lua_value(cur_col, x) for x in tmp))
                one_row.append("}")
            if j != len(col_desc_list) - 1:
                one_row.append(",")
        one_row.append("},")
        sb.append("".join(one_row))
        sb.append("\n")
    sb.append("}\n")

    helperstr = "\tif helper[b] then\n" "\t\treturn helper[b]\n" "\tend\n"
    substr = (
        "local mt{0} = {{}}\n"
        "mt{0}.__index = function(a,b)\n"
        "\tif fieldIdx{0}[b] then\n"
        "\t\treturn a[fieldIdx{0}[b]]\n"
        "\tend\n{1}"
        "\treturn nil\n"
        "end\n"
        "mt{0}.__metatable = false\n"
        "for _,v in pairs(data) do\n"
    )

    sb.append(f"local helper = jy_utils.prequire('Jyx2Configs/{config_type}Helper')\n")
    sb.append(substr.format("", helperstr))
    sb.append("\tsetmetatable(v,mt)\nend\n")

    for col in col_desc_list:
        if col.type == L_STRUCT and len(col.sub_name) > 1:
            sb.append(f"local fieldIdx{col.name} = {{}}\n")
            for j, sub in enumerate(col.sub_name):
                sb.append(f"fieldIdx{col.name}.{sub} = {j + 1}\n")
            sb.append(substr.format(col.name, ""))
            sb.append(
                f"\tfor _,t in pairs(v.{col.name}) do\n"
                "\t\tif type(t) == 'table' then\n"
                f"\t\t\tsetmetatable(t,mt{col.name})\n"
                "\t\tend\n"
                "\tend\nend\n"
            )

    sb.append("local configMgr = Jyx2:GetModule('ConfigMgr')\n")
    sb.append(f"configMgr:AddConfigTable([[{config_type}]], data)")
    return "".join(sb)


def read_sheet(path: str) -> tuple[list[list[str]], int, int]:
    """读取第一张工作表，返回 (rows, colNum, rowNum)，单元格都转成字符串"""
    if openpyxl is None:
        raise RuntimeError("需要 openpyxl：python -m pip install openpyxl")
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        ws = wb.worksheets[0]
        col_num = ws.max_column or 0
        row_num = ws.max_row or 0
        rows: list[list[str]] = []
        for r in ws.iter_rows(values_only=True):
            row = [cell_to_str(v) for v in r]
            if len(row) < col_num:
                row.extend([""] * (col_num - len(row)))
            rows.append(row)
        while len(rows) < row_num:
            rows.append([""] * col_num)
        return rows, col_num, row_num
    finally:
        wb.close()


def export_single_lua_file(path: str, out_dir: str) -> str | None:
    """对应 C# ExportSingleLuaFile：返回生成的文件路径，未生成返回 None"""
    rows, col_num, row_num = read_sheet(path)
    if not rows or col_num < 2 or rows[0][0] != "LuaConfigGen":
        return None
    out_path = os.path.join(out_dir, rows[0][1] + "Config.lua")
    content = gen_lua_file(rows, col_num, row_num)
    if content:
        os.makedirs(out_dir, exist_ok=True)
        # C# 的 File.WriteAllText 默认 UTF-8（无 BOM）
        with open(out_path, "w", encoding="utf-8", newline="\n") as f:
            f.write(content)
        return out_path
    return None


def export_all_lua_file(in_dir: str, out_dir: str, clean: bool = True) -> list[str]:
    """对应 C# ExportAllLuaFile：批量导出（clean=True 时先清空输出目录）"""
    if clean and os.path.isdir(out_dir):
        import shutil

        shutil.rmtree(out_dir)
    os.makedirs(out_dir, exist_ok=True)

    results: list[str] = []
    for root, _dirs, files in os.walk(in_dir):
        for name in files:
            if not name.lower().endswith(".xlsx") or "~$" in name:
                continue
            out = export_single_lua_file(os.path.join(root, name), out_dir)
            if out:
                results.append(out)
    return results


def main(argv: list[str] | None = None) -> int:
    import argparse

    parser = argparse.ArgumentParser(description="Excel 配置表 -> Lua（ExcelToLua.cs 的 CLI 版本）")
    parser.add_argument("in_dir", help="Excel 目录，例如 jyx2/Assets/Mods/SAMPLE/Configs")
    parser.add_argument("-o", "--out-dir", default=None, help="输出目录，默认 <in_dir>/Lua")
    parser.add_argument("--no-clean", action="store_true", help="不清空输出目录")
    args = parser.parse_args(argv)

    out_dir = args.out_dir or os.path.join(args.in_dir, "Lua")
    files = export_all_lua_file(args.in_dir, out_dir, clean=not args.no_clean)
    for f in files:
        print(f"[ok] {os.path.basename(f)}")
    print(f"共生成 {len(files)} 个 Lua 配置表 -> {out_dir}")
    return 0 if files else 1


if __name__ == "__main__":
    raise SystemExit(main())

